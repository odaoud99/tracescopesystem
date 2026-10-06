from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DC
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).parent
NAVY='0B121B'; PANEL='121F2E'; PANEL2='172638'; LINE='2B3D50'; TEXT='E7EEF6'; MUTED='9AAABD'; CYAN='72D8E7'; GREEN='78E3A4'; GOLD='E5BA66'; LILAC='B6A3FF'
def col(s): return RGBColor.from_string(s)

# Presentation: concise product-demo arc, widescreen 16:9.
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5);blank=p.slide_layouts[6]
def box(s,x,y,w,h,fill=PANEL,line=LINE,kind=MSO_SHAPE.ROUNDED_RECTANGLE):
 z=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));z.fill.solid();z.fill.fore_color.rgb=col(fill);z.line.color.rgb=col(line);z.line.width=Pt(1);return z
def text(s,x,y,w,h,t,size=14,color=TEXT,bold=False,align=None):
 z=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=z.text_frame;f.clear();f.word_wrap=True;f.margin_left=Inches(.04);f.margin_right=Inches(.04);f.margin_top=Inches(.03)
 for i,line in enumerate(str(t).split('\n')):
  q=f.paragraphs[0] if i==0 else f.add_paragraph();q.text=line;q.font.name='Aptos';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=col(color);q.space_after=Pt(6)
  if align is not None:q.alignment=align
 return z
def base(title,n):
 s=p.slides.add_slide(blank);s.background.fill.solid();s.background.fill.fore_color.rgb=col(NAVY);box(s,0,0,.1,7.5,GREEN,GREEN,MSO_SHAPE.RECTANGLE);text(s,.55,.24,11,.24,'TRACESCOPE  /  CORE NETWORK TRACE WORKBENCH',9,GREEN,True)
 if title:text(s,.55,.68,12.1,.6,title,26,TEXT,True)
 box(s,.55,7.08,12.2,.012,LINE,LINE,MSO_SHAPE.RECTANGLE);text(s,.58,7.15,11,.18,'ENGINEER PRODUCT OVERVIEW  ·  05 OCT 2026',8,MUTED);text(s,12.12,7.13,.5,.22,f'{n:02d}',9,MUTED,True,PP_ALIGN.RIGHT);return s
def card(s,x,y,w,h,head,body,accent=CYAN):
 box(s,x,y,w,h);box(s,x,y,.05,h,accent,accent,MSO_SHAPE.RECTANGLE);text(s,x+.18,y+.15,w-.36,.32,head,14,accent,True);text(s,x+.18,y+.57,w-.36,h-.68,body,11,TEXT)
def node(s,x,y,label,sub,accent):
 box(s,x,y,1.62,.68,PANEL2,accent);text(s,x+.05,y+.1,1.52,.22,label,11,TEXT,True,PP_ALIGN.CENTER);text(s,x+.06,y+.38,1.5,.18,sub,8,MUTED,False,PP_ALIGN.CENTER)

# 1 cover
s=p.slides.add_slide(blank);s.background.fill.solid();s.background.fill.fore_color.rgb=col(NAVY);box(s,.7,.75,.08,5.9,GREEN,GREEN,MSO_SHAPE.RECTANGLE)
text(s,1.05,1,10,.3,'TELECOM TRACE ANALYSIS  /  PRODUCT BRIEF',10,GREEN,True);text(s,1.04,1.6,10,1,'TraceScope',43,TEXT,True);text(s,1.08,2.72,10,.8,'From raw network records to an engineer-readable subscriber journey.',22,MUTED)
text(s,1.08,3.8,9.6,.72,'A local-first workbench for LMT user traces, protocol context, and technology-aware network paths.',15,TEXT)
for i,(a,b,c) in enumerate([('UE','subscriber',CYAN),('RAN','radio access',GREEN),('CORE','control',GOLD),('DATA','service',LILAC)]):
 x=1.08+i*2.8;node(s,x,5.45,a,b,c)
 if i<3:text(s,x+1.8,5.6,.65,.3,'→',19,GREEN,True,PP_ALIGN.CENTER)

# 2 problem/solution
s=base('The trace is evidence. The engineer needs the story.',2)
card(s,.7,1.7,3.8,2.2,'RAW EXPORT','Binary containers, internal labels, protocol messages, and node positions arrive together.',CYAN)
card(s,4.76,1.7,3.8,2.2,'ENGINEERING FRICTION','Finding the right event and following signaling across technologies takes repeated manual inspection.',GOLD)
card(s,8.82,1.7,3.8,2.2,'TRACE SCOPE','A guided timeline, message inspector, call flow, and infrastructure route keep evidence in one review surface.',GREEN)
text(s,.82,4.45,11.6,.45,'Make each record easier to locate, read, and relate to its neighbors.',20,TEXT,True)
text(s,.84,5.15,11.2,.85,'TraceScope is an analysis aid—not a full vendor dissector or live network inventory. It separates values found in the file from protocol-based interpretation.',13,MUTED)

# 3 workflow
s=base('One workflow, from import to reviewable evidence',3)
steps=[('01 · IMPORT','Open PTMF or supported text trace. Parsing runs in the browser.'),('02 · MAP','Inspect time, position, message type, direction, and payload boundary.'),('03 · FOLLOW','Move through timeline, sequence flow, and technology-aware route.'),('04 · EXPLAIN','Read telecom context, consult references, and capture findings.')]
for i,(a,b) in enumerate(steps):
 x=.7+i*3.12;card(s,x,1.8,2.68,3.65,a,b,[CYAN,GREEN,GOLD,LILAC][i])
 if i<3:text(s,x+2.7,3.25,.38,.3,'→',19,GREEN,True,PP_ALIGN.CENTER)
text(s,.83,5.9,11,.35,'Unknown fields remain unknown; no successful procedure is inferred from a request alone.',12,MUTED)

# 4 viewer
s=base('A message inspector built for telecom context',4)
box(s,.7,1.55,7.5,4.95,'0E1925');text(s,.95,1.82,5,.25,'MESSAGE TIMELINE',10,CYAN,True)
rows=[('16:13:04.395','SGsAP-PAGING-REQUEST','MSC → MME',GOLD),('16:13:04.402','S1AP · PAGING','MME → eNodeB',CYAN),('16:13:04.671','EXTENDED SERVICE REQUEST','UE → MME',GREEN),('16:13:04.672','MM_SGSAP_SERVICE_REQ','Internal · vendor label',LILAC)]
for i,(a,b,c,d) in enumerate(rows):
 y=2.25+i*.83;box(s,.98,y,6.9,.63,PANEL,LINE);text(s,1.13,y+.1,1.45,.22,a,9,MUTED);text(s,2.62,y+.08,4.9,.24,b,10,TEXT,True);text(s,2.62,y+.34,4.9,.18,c,8,d)
card(s,8.55,1.55,4.05,2.15,'OBSERVED FIELDS','Timestamp · node/LMT · direction if decoded · property · type/code · size · offset · bytes',CYAN)
card(s,8.55,3.95,4.05,2.55,'TELECOM KNOWLEDGE','Plain-language message role, basis, and 3GPP reference appear beside the selected record. Explanations state what the message alone cannot prove.',GREEN)

# 5 technology routes
s=base('Follow the route across the detected technology',5)
data=[('2G','MS','BTS','BSC / PCU','MSC / SGSN','GGSN',CYAN),('3G','UE','Node B','RNC','SGSN','GGSN',GREEN),('4G','UE','eNodeB','MME / SGW','PGW','PDN',GOLD),('5G','UE','gNodeB','AMF / SMF','UPF','DN',LILAC)]
for i,row in enumerate(data):
 y=1.55+i*1.02;text(s,.72,y+.22,.7,.28,row[0],13,row[6],True)
 for j,lab in enumerate(row[1:6]):
  x=1.55+j*2.16;node(s,x,y,lab,'reference component',row[6])
  if j<4:text(s,x+1.7,y+.18,.4,.25,'→',16,row[6],True,PP_ALIGN.CENTER)
text(s,.8,5.94,5.1,.28,'Green = observed match',10,GREEN,True);text(s,5.05,5.94,6.8,.28,'Gold dashed links in the app = inferred standard topology',10,GOLD,True)

# 6 evidence boundaries
s=base('Keep the evidence boundary visible',6)
card(s,.7,1.7,3.8,3.7,'OBSERVED','Values read from record containers or explicit endpoints: timestamp, position, type/code, length, and available direction.',GREEN)
card(s,4.76,1.7,3.8,3.7,'MAPPED','Known signatures associate a record with a likely technology or route segment. The UI labels these associations as observed or inferred.',GOLD)
card(s,8.82,1.7,3.8,3.7,'EXPLAINED','Message knowledge describes the protocol role and states when a response, cause code, or decoded information element is needed.',CYAN)
text(s,.83,5.87,11,.4,'A request is not proof of success. A vendor-internal label is not a decoded information element.',14,TEXT,True)

# 7 companion features
s=base('Engineer tools with explicit data boundaries',7)
card(s,.7,1.65,3.8,3.85,'AI ASSISTANT','Text and browser-supported voice\nTrace-grounded questions with record citations\nLocal viewer commands stay local\nNVIDIA GLM-5.3 via local server',CYAN)
card(s,4.76,1.65,3.8,3.85,'RESEARCH DESK','Explicit Huawei-focused Tavily search\nHuawei-domain filter\nSaved sources and documents in browser storage',GREEN)
card(s,8.82,1.65,3.8,3.85,'AUDIT & EXPORT','Activities separated from why/decisions\nBrowser-local audit trail\nTrace and audit JSON export',LILAC)
text(s,.84,5.92,11,.32,'Selected trace context is sent to NVIDIA when asking AI questions. Web research is user-triggered.',11,GOLD,True)

# 8 walkthrough and limit
s=base('A practical engineer walkthrough',8)
walk=[('1','Import the trace','Confirm record count, timestamp range, and detected technologies.'),('2','Locate a message','Search by record number, protocol, type, or node.'),('3','Inspect the evidence','Read fields, telecom information, basis, and raw bytes.'),('4','Follow the procedure','Open Call flow or Network path; step or animate.'),('5','Capture findings','Ask a question, check the audit trail, export JSON.')]
for i,(n,a,b) in enumerate(walk):
 y=1.55+i*.83;box(s,.85,y,.48,.45,PANEL2,GREEN,MSO_SHAPE.OVAL);text(s,.85,y+.08,.48,.23,n,10,GREEN,True,PP_ALIGN.CENTER);text(s,1.55,y,2.55,.3,a,13,TEXT,True);text(s,4.18,y,7.8,.42,b,11,MUTED)
text(s,.9,6.05,11,.42,'Validate conclusions with a qualified dissector, applicable 3GPP release, and network-side evidence.',12,GOLD,True)

p.save(OUT/'TraceScope_Project_Presentation.pptx')

# User manual (DOCX)
d=Document();sec=d.sections[0];sec.top_margin=DI(.7);sec.bottom_margin=DI(.65);sec.left_margin=DI(.8);sec.right_margin=DI(.8)
styles=d.styles;styles['Normal'].font.name='Aptos';styles['Normal'].font.size=DP(10);styles['Normal'].font.color.rgb=DC(40,53,68)
for n,z,c in [('Title',31,'102334'),('Heading 1',20,'102334'),('Heading 2',14,'17645A')]:styles[n].font.name='Aptos Display';styles[n].font.size=DP(z);styles[n].font.bold=True;styles[n].font.color.rgb=DC.from_string(c)
sec.header.paragraphs[0].text='TRACESCOPE  /  ENGINEER FIELD GUIDE';sec.header.paragraphs[0].runs[0].font.color.rgb=DC(23,100,90)
sec.footer.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT;sec.footer.paragraphs[0].text='TraceScope · User manual · 05 Oct 2026'
def h(t,l=1):d.add_heading(t,level=l)
def para(t):d.add_paragraph(t)
def bullets(xs):
 for x in xs:d.add_paragraph(x,style='List Bullet')
def nums(xs):
 for x in xs:d.add_paragraph(x,style='List Number')
def tbl(heads,rows):
 t=d.add_table(rows=1,cols=len(heads));t.style='Light Shading Accent 1';
 for c,v in zip(t.rows[0].cells,heads):c.text=v
 for c in t.rows[0].cells:
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173C45');c._tc.get_or_add_tcPr().append(sh)
  for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=DC(255,255,255)
 for row in rows:
  cs=t.add_row().cells
  for c,v in zip(cs,row):c.text=v
 return t
d.add_paragraph('TRACESCOPE',style='Title');d.add_paragraph('User manual & engineer walkthrough',style='Subtitle');para('A practical guide to importing, reading, and explaining mobile network user traces.')
call=d.add_table(rows=1,cols=1);call.cell(0,0).text='SCOPE  This guide describes the current local TraceScope implementation. Distinguish fields observed in the file from protocol interpretations and topology links inferred from standards.'
d.add_page_break()
h('1. Project description');para('TraceScope is a browser-based analysis workbench for Huawei LMT user-trace exports and supported text traces. It organizes records into a searchable timeline, call-flow sequence, technology-aware infrastructure route, and message inspector with telecom context.')
h('Who it is for',2);bullets(['Core-network and mobile-access engineers reviewing paging, mobility, service, and session signaling.','Operations/support teams preparing evidence-backed trace summaries.','Telecom learners who need message and network-component context.'])
h('Main capabilities',2);tbl(['Area','Purpose'],[('Trace import','Parse PTMF record structures and supported text/log rows in the browser.'),('Timeline & inspector','Review time, position, message fields, decoded direction when available, payload length, and bytes.'),('Call flow','Show records against participant lanes; arrows appear only when endpoints are decoded.'),('Network path','Vendor-neutral 2G/3G/4G/5G reference topology with detected-technology selection and record stepping.'),('Telecom knowledge','Plain-language message descriptions, interpretation basis, and 3GPP catalogue links.'),('Assistant','Text/voice questions via NVIDIA GLM-5.3; answers can cite records.'),('Research & audit','Tavily searches, browser-local references, Activities and Why/decisions views.'),('Export','Trace analysis and audit JSON downloads.')])
h('2. Start TraceScope');nums(['Open the project folder in a terminal.','Start the local Node.js server with `node server.mjs`.','Open http://127.0.0.1:4173/ in a supported browser.','For AI and web research, set NVIDIA_API_KEY and TAVILY_API_KEY in the ignored .env.local file, then restart the server. Never put credentials in browser code.'])
para('The browser parses the chosen trace. The local Node server serves the interface and proxies external API calls. TraceScope is bound to loopback for the local development workflow.')
h('Data handling',2);para('Each non-command assistant question can send the question and a compact set of relevant trace fields to NVIDIA. Subscriber identifiers are omitted unless explicitly requested; raw payload bytes are excluded. Tavily is called only when a user submits a search. Audit entries and saved research documents remain in the current browser profile. Follow organizational trace-handling policy.')
h('3. Import and read a trace');nums(['Choose Import trace / Choose trace file and select a PTMF, BIN, DAT, TRACE, TXT, or LOG file.','Use the sample option to explore the supplied example.','Confirm message count, node count, duration, format, and timestamps after import.','Search for a message type, protocol, node name, or text; select a row to open the inspector.'])
h('The parser may show',2);bullets(['Record boundaries, offsets/lengths, valid timestamps, and printable strings.','A known protocol/type label when its identifier is recognized.','Source/destination and direction only when present in decoded fields or supported text rows.','Raw bytes for protocol fields that have not been decoded.'])
para('“Not decoded” means the current parser did not establish that field from this export. It does not mean the network omitted it.')
h('4. Inspect a message');tbl(['Inspector area','Use'],[('Record fields','Review time, position/LMT, direction, property, type, code, size, offset, identifier field, and technology evidence.'),('Telecom information','Read the message role, basis, and available standards link.'),('Observed text','Inspect printable strings found in the record span.'),('Hex','Review the first displayed payload bytes.'),('How to read','Review parser boundaries and what needs a deeper dissector.')])
para('A request records an attempted procedure step. Inspect later responses, reject causes, and completion messages before concluding that paging, registration, attach, or service setup succeeded.')
h('5. Call flow and Network path');h('Call flow',2);bullets(['Participant lifelines show messages in chronological order.','Use filters and Show all records to navigate longer traces.','Arrows are drawn only for decoded source/destination endpoints; otherwise, the message remains an observation at its known LMT/node lane.','Select a row to open that message in the inspector.'])
h('Network path',2);bullets(['AUTO follows the technology associated with the selected record, or the strongest detected technology when none is selected.','Choose a generation filter, or ALL MESSAGES when multiple generations were detected.','Use Previous, Next, and Animate route to step through matching messages.','Select a message to highlight its matched components and see interpretation and standards references.','Green = position/endpoint observed. Gold dashed connectors = protocol-associated or standard reference links. Neutral nodes = reference topology only.','Unmapped internal messages remain visible in the list without a guessed hop.'])
para('The topology is vendor-neutral and synthesized from the supplied infrastructure poster and standards knowledge. It is not a carrier inventory and does not prove that every displayed link was used in the capture.')
h('6. Chat and voice assistant');para('Open Chat with trace or Voice assistant. Example questions: “Summarize this trace”, “What happened at record 13?”, “Why was 4G detected?”, or “Explain the Extended Service Request and what confirms success?” Record citations link back to trace events.')
bullets(['Viewer commands such as opening tabs, selecting records, playback, and export are handled locally.','The assistant’s knowledge answer uses the configured NVIDIA GLM-5.3 endpoint.','Voice recognition/output depends on browser support and may use the browser’s speech service.','Do not send trace data to an external model unless organizational policy permits it.'])
h('7. Research Desk');nums(['Open Research desk.','Enter a Huawei/LMT/protocol query and optionally restrict results to Huawei domains.','Submit Search explicitly; a search may consume a Tavily credit.','Review the source and save useful items to the local library.','Add PDF, Word, text, CSV, HTML, or JSON reference documents to browser storage.'])
para('Web search results are external reference material, not evidence from the trace. Confirm version, release, and applicability in the source document.')
h('8. Audit and export');bullets(['Activities records UI actions such as imports, selections, research searches, and exports.','Why / decisions displays the observed evidence and rules behind local analysis—not hidden chain-of-thought.','Export the visible audit tab as JSON. Export report downloads trace metadata and message-level analysis as JSON.','Treat exported trace reports as sensitive data.'])
h('9. Evidence levels and limitations');tbl(['Level','Meaning'],[('Observed','Value appears explicitly in the record/container or decoded endpoint.'),('Translated','Protocol/type identifier recognized; not necessarily all information elements.'),('Associated','Protocol signature points to likely network roles; exact endpoints may not be decoded.'),('Reference','Generic 3GPP-style topology for explanation, not the operator live route.'),('Unknown','Evidence is insufficient; do not fill the gap by guessing.')])
bullets(['Binary PTMF decoding is partial; some vendor-internal records remain labels and raw bytes.','Technology recognition depends on exported protocol/node signatures. Missing signatures do not prove a technology was unused.','Route view is a reference model, not an end-to-end topology reconstruction.','Assistant responses may be incomplete or inaccurate; validate using a qualified dissector, applicable 3GPP release, and network-side evidence.'])
h('10. Troubleshooting');tbl(['Issue','Action'],[('Trace will not load','Confirm file format; try the sample or a text export if binary record markers are unfamiliar.'),('Technology not confirmed','Check the translated type and evidence list; the view needs a recognized RAT-specific signature.'),('No route highlight','Select a message with a known endpoint/protocol. Vendor-internal positions may not map to a generic role.'),('AI unavailable','Check NVIDIA_API_KEY in .env.local, restart the server, and inspect /api/health. Never share the key.'),('Research unavailable','Check TAVILY_API_KEY and server status; submit a query from Research Desk.'),('Saved items missing','They belong to the same browser profile/site storage; clearing site data removes them.')])
h('Presenter walkthrough');nums(['Import the sample and point out record count and timestamp range.','Select a translated message; show observed fields, interpretation basis, and 3GPP link.','Open Call flow and explain why undecoded endpoints do not get arrows.','Open Network path in AUTO and animate a few events; distinguish green matches from inferred links.','Ask an assistant question and open its record citation; explain NVIDIA context sharing.','Finish with Activities / Why decisions and export a JSON report.'])
d.save(OUT/'TraceScope_User_Manual.docx')

description='''# TraceScope — Project description\n\n**TraceScope** is a local-first telecom trace analysis workbench for core-network engineers. It organizes Huawei LMT user-trace exports and supported text logs into a searchable event timeline, call-flow sequence, technology-aware infrastructure route, and message inspector with telecom context.\n\n## Purpose\n\nMake exported signaling easier to review one record at a time: identify what the file contains, which protocol label was recognized, what role a message commonly plays, and which evidence remains missing.\n\n## Capabilities\n\n- Browser-side parsing of Huawei PTMF structures and supported text/log rows.\n- Timeline search and message inspection with timestamps, positions, decoded fields, printable strings, and a bounded hex view.\n- Call-flow sequence with participant lanes and arrows only for decoded endpoints.\n- Vendor-neutral 2G GSM/GPRS, 3G UMTS, 4G LTE/EPC, and 5G 5GS route references that follow trace technology signatures.\n- Telecom message explanations, interpretation basis, and 3GPP specification links.\n- Text/voice assistant via NVIDIA GLM-5.3, with record citations and local viewer commands.\n- Explicit Tavily Huawei research, browser-local documents/sources, activity/decision audit, and JSON export.\n\n## Operating model\n\nThe browser parses the selected file. A local Node.js server serves the UI and proxies external APIs. Credentials belong in the ignored `.env.local` file and are not delivered to browser code. The app is served at `http://127.0.0.1:4173/`.\n\n## Evidence and limitations\n\nTraceScope is an engineering review aid, not a full vendor protocol dissector or live network inventory. It recognizes a subset of Huawei PTMF message identifiers and protocol signatures. A decoded label is not proof that every information element was interpreted. Dashed topology links are reference/inferred relationships, not proof of a hop in the capture. A request does not establish successful procedure completion. Validate conclusions against raw records, a qualified dissector, the applicable 3GPP release, and network-side evidence.\n\nWhen configured, assistant questions send selected trace context to NVIDIA; subscriber identifiers are omitted by default and raw payload bytes are excluded. Research queries go to Tavily only after explicit submission. Audit and saved research items are stored in the current browser profile. Follow organizational trace-handling policy.\n\n## Audience\n\nMobile core/access engineers, network operations and support teams, and telecom trainees.\n\n## Current implementation\n\nSingle-page browser UI (`index.html`), local Node.js server (`server.mjs`), and network path module (`public/network-path.js`).\n'''
(OUT/'TraceScope_Project_Description.md').write_text(description,encoding='utf-8')
print('Created presentation, user manual, and project description in',OUT)
